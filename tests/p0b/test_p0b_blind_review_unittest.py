from __future__ import annotations

import csv
import json
import unittest
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image

from p0b_deca.blind_review import SOURCE_MODES, _leakage_audit, sha256


ROOT = Path("data/processed/P0B_DECA_Pilot12_v1/blind_review_v1")


class BlindReviewPackageTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        if not ROOT.is_dir():
            raise unittest.SkipTest("blind review package has not been built")
        cls.reviewer = ROOT / "reviewer_package"
        cls.private = ROOT / "private_unblinding"
        with (cls.private / "unblinding_key.csv").open(newline="", encoding="utf-8-sig") as handle:
            cls.key = list(csv.DictReader(handle))

    def test_complete_unique_balanced_mapping(self) -> None:
        self.assertEqual(len(self.key), 12)
        self.assertEqual({row["blind_case_id"] for row in self.key}, {f"BR{i:02d}" for i in range(1, 13)})
        self.assertEqual(sum(row["mode_A_true_name"] == SOURCE_MODES[0] for row in self.key), 6)
        self.assertEqual(sum(row["mode_A_true_name"] == SOURCE_MODES[1] for row in self.key), 6)
        for row in self.key:
            self.assertEqual({row["mode_A_true_name"], row["mode_B_true_name"]}, set(SOURCE_MODES))

    def test_panels_and_zip_are_reviewer_only(self) -> None:
        for number in range(1, 13):
            for kind in ("overview", "relighting"):
                path = self.reviewer / "panels" / f"BR{number:02d}_{kind}.png"
                self.assertTrue(path.is_file())
                with Image.open(path) as image:
                    self.assertFalse(image.info)
        package = ROOT / "P0B_blind_review_reviewer_package.zip"
        if not package.is_file():
            self.skipTest("historical reviewer ZIP is absent; package content cannot be verified")
        with zipfile.ZipFile(package) as archive:
            names = archive.namelist()
        self.assertEqual(len(names), len([path for path in self.reviewer.rglob("*") if path.is_file()]))
        self.assertTrue(all(name.startswith("reviewer_package/") for name in names))
        self.assertFalse(any("private_unblinding" in name for name in names))

    def test_reviewer_package_has_no_mapping_leakage(self) -> None:
        forbidden = set(SOURCE_MODES) | {"Control", "Patient", "Xiaomi", "HONOR", "camera_make", "camera_model", "binary_name"}
        forbidden.update(row["original_case_id"] for row in self.key)
        self.assertEqual(_leakage_audit(self.reviewer, forbidden), [])

    def test_workbook_has_12_rows_validations_and_summary_formulas(self) -> None:
        workbook = load_workbook(self.reviewer / "review_form/P0B_input_mode_blind_review_template.xlsx", data_only=False)
        self.assertEqual(workbook.sheetnames, ["Instructions", "Case_Review", "Lists", "Summary", "Codebook"])
        review = workbook["Case_Review"]
        self.assertEqual([review.cell(row, 1).value for row in range(2, 14)], [f"BR{i:02d}" for i in range(1, 13)])
        self.assertEqual(len(review.data_validations.dataValidation), 5)
        self.assertTrue(any(str(cell.value).startswith("=") for row in workbook["Summary"].iter_rows() for cell in row if isinstance(cell.value, str)))
        self.assertEqual(workbook["Lists"].sheet_state, "hidden")

    def test_key_and_hashes_recover_all_sources_without_changes(self) -> None:
        with (self.private / "source_file_hashes.csv").open(newline="", encoding="utf-8-sig") as handle:
            hashes = list(csv.DictReader(handle))
        source_hashes = [row for row in hashes if row["scope"] == "source_result"]
        self.assertEqual(len(source_hashes), 24 * 14)
        for row in source_hashes:
            self.assertEqual(sha256(Path(row["path"])), row["sha256"])
        self.assertEqual(len({row["original_manifest_row"] for row in self.key}), 12)
        self.assertEqual(json.loads((self.private / "build_manifest.json").read_text(encoding="utf-8"))["no_deca_inference_run"], True)


if __name__ == "__main__":
    unittest.main()
