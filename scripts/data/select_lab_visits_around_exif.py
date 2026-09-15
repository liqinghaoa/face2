"""Select laboratory records from hospital visits surrounding image capture.

For each image sample, a terminal numeric suffix (for example ``-1``) is
removed to obtain the hospital patient ID. Hospital visits are anchored by the
earliest laboratory application time within each visit. When both sides exist,
the nearest visit before and after image capture are selected. If either side
is absent, the two visits with the smallest absolute time difference are used.
If only one visit exists, that visit is retained.
"""

from __future__ import annotations

import argparse
import csv
import json
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_SAMPLE_CSV = ROOT / "data/processed/global_face_R3DPR/nyha_2class_with_exif_time.csv"
DEFAULT_LAB_XLSX = ROOT / "data/raw/table/心功能患者数据-检验数据.xlsx"
DEFAULT_DETAIL_JSONL = ROOT / ".codex_tmp/selected_lab_two_visits.jsonl"
DEFAULT_AUDIT_JSON = ROOT / ".codex_tmp/selected_lab_two_visits_audit.json"


def parse_datetime(value: object) -> datetime:
    if isinstance(value, datetime):
        return value
    if value is None or not str(value).strip():
        raise ValueError("Blank datetime")
    text = str(value).strip().replace(":", "-", 2)
    return datetime.strptime(text, "%Y-%m-%d %H:%M:%S")


def json_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    return value


def patient_id(sample_id: str) -> str:
    return re.sub(r"-\d+$", "", sample_id)


def relative_position(anchor: datetime, capture: datetime) -> str:
    if anchor < capture:
        return "EXIF前"
    if anchor > capture:
        return "EXIF后"
    return "与EXIF同时"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--sample-csv", type=Path, default=DEFAULT_SAMPLE_CSV)
    parser.add_argument("--lab-xlsx", type=Path, default=DEFAULT_LAB_XLSX)
    parser.add_argument("--detail-jsonl", type=Path, default=DEFAULT_DETAIL_JSONL)
    parser.add_argument("--audit-json", type=Path, default=DEFAULT_AUDIT_JSON)
    args = parser.parse_args()

    with args.sample_csv.open(encoding="utf-8-sig", newline="") as handle:
        sample_rows = list(csv.DictReader(handle))
    if not sample_rows or "ID" not in sample_rows[0] or "EXIF_DateTimeOriginal" not in sample_rows[0]:
        raise ValueError("Sample CSV must contain ID and EXIF_DateTimeOriginal columns")
    sample_ids = [row["ID"].strip() for row in sample_rows]
    if len(set(sample_ids)) != len(sample_ids):
        raise ValueError("Sample IDs must be unique")

    workbook = load_workbook(args.lab_xlsx, read_only=True, data_only=True)
    worksheet = workbook.worksheets[0]
    raw_headers = list(next(worksheet.iter_rows(values_only=True)))
    if len(raw_headers) < 7:
        raise ValueError("Laboratory workbook has fewer columns than expected")
    source_headers = ["源表序号" if value is None else str(value) for value in raw_headers]

    # First pass: summarize each admission without retaining all 270k rows.
    visits: dict[str, dict[str, dict[str, Any]]] = defaultdict(dict)
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        pid = str(row[1]).strip() if row[1] is not None else ""
        visit_no = str(row[2]).strip() if row[2] is not None else ""
        if not pid or not visit_no:
            continue
        application_time = parse_datetime(row[5])
        visit = visits[pid].setdefault(
            visit_no,
            {"anchor": application_time, "last": application_time, "row_count": 0},
        )
        visit["anchor"] = min(visit["anchor"], application_time)
        visit["last"] = max(visit["last"], application_time)
        visit["row_count"] += 1

    selected_by_visit: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    audit_rows: list[dict[str, Any]] = []
    rule_counts: dict[str, int] = defaultdict(int)
    selected_visit_total = 0

    for sample in sample_rows:
        sample_id = sample["ID"].strip()
        pid = patient_id(sample_id)
        capture = parse_datetime(sample["EXIF_DateTimeOriginal"])
        candidates = [
            {"visit_no": visit_no, **summary}
            for visit_no, summary in visits.get(pid, {}).items()
        ]
        if not candidates:
            raise ValueError(f"No laboratory admissions found for sample {sample_id}")

        before = [visit for visit in candidates if visit["anchor"] <= capture]
        after = [visit for visit in candidates if visit["anchor"] > capture]
        if before and after:
            chosen = [
                max(before, key=lambda visit: visit["anchor"]),
                min(after, key=lambda visit: visit["anchor"]),
            ]
            rule = "EXIF前后各最近一次"
            labels = ["EXIF前最近", "EXIF后最近"]
        else:
            chosen = sorted(
                candidates,
                key=lambda visit: (
                    abs((visit["anchor"] - capture).total_seconds()),
                    visit["anchor"],
                    visit["visit_no"],
                ),
            )[:2]
            if len(chosen) == 1:
                rule = "仅一次住院可用"
                labels = ["唯一可用住院"]
            else:
                rule = "缺少一侧，选择距离最近两次"
                labels = ["距离最近1", "距离最近2"]

        rule_counts[rule] += 1
        selected_visit_total += len(chosen)
        selected_audit: list[dict[str, Any]] = []
        for rank, (visit, label) in enumerate(zip(chosen, labels), start=1):
            gap_days = abs((visit["anchor"] - capture).total_seconds()) / 86400
            descriptor = {
                "sample_id": sample_id,
                "patient_id": pid,
                "capture": capture,
                "rule": rule,
                "rank": rank,
                "label": label,
                "position": relative_position(visit["anchor"], capture),
                "anchor": visit["anchor"],
                "gap_days": round(gap_days, 4),
            }
            selected_by_visit[(pid, visit["visit_no"])].append(descriptor)
            selected_audit.append({**descriptor, **visit})

        audit: dict[str, Any] = {
            "样本ID": sample_id,
            "患者ID": pid,
            "EXIF原始拍摄时间": capture.strftime("%Y-%m-%d %H:%M:%S"),
            "可用住院次数": len(candidates),
            "选择规则": rule,
            "选中住院数": len(chosen),
        }
        for index in range(2):
            prefix = f"住院{index + 1}"
            if index < len(selected_audit):
                item = selected_audit[index]
                audit.update(
                    {
                        f"{prefix}次数": item["visit_no"],
                        f"{prefix}选择标签": item["label"],
                        f"{prefix}相对位置": item["position"],
                        f"{prefix}锚点申请时间": item["anchor"].strftime("%Y-%m-%d %H:%M:%S"),
                        f"{prefix}最晚申请时间": item["last"].strftime("%Y-%m-%d %H:%M:%S"),
                        f"{prefix}时间差_天": item["gap_days"],
                        f"{prefix}检验行数": item["row_count"],
                    }
                )
            else:
                audit.update(
                    {
                        f"{prefix}次数": None,
                        f"{prefix}选择标签": None,
                        f"{prefix}相对位置": None,
                        f"{prefix}锚点申请时间": None,
                        f"{prefix}最晚申请时间": None,
                        f"{prefix}时间差_天": None,
                        f"{prefix}检验行数": None,
                    }
                )
        audit_rows.append(audit)

    detail_headers = [
        "样本ID",
        "患者ID",
        "EXIF原始拍摄时间",
        "选择规则",
        "选择顺序",
        "选择标签",
        "相对EXIF位置",
        "住院锚点申请时间",
        "时间差_天",
        "源表行号",
        *source_headers,
    ]
    args.detail_jsonl.parent.mkdir(parents=True, exist_ok=True)
    detail_count = 0
    with args.detail_jsonl.open("w", encoding="utf-8", newline="\n") as output:
        output.write(json.dumps({"headers": detail_headers}, ensure_ascii=False) + "\n")
        # Second pass: stream only rows belonging to selected admissions.
        worksheet = load_workbook(args.lab_xlsx, read_only=True, data_only=True).worksheets[0]
        next(worksheet.iter_rows(values_only=True))
        for source_row, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
            pid = str(row[1]).strip() if row[1] is not None else ""
            visit_no = str(row[2]).strip() if row[2] is not None else ""
            for descriptor in selected_by_visit.get((pid, visit_no), []):
                values = [
                    descriptor["sample_id"],
                    descriptor["patient_id"],
                    descriptor["capture"].strftime("%Y-%m-%d %H:%M:%S"),
                    descriptor["rule"],
                    descriptor["rank"],
                    descriptor["label"],
                    descriptor["position"],
                    descriptor["anchor"].strftime("%Y-%m-%d %H:%M:%S"),
                    descriptor["gap_days"],
                    source_row,
                    *[json_value(value) for value in row],
                ]
                output.write(json.dumps(values, ensure_ascii=False) + "\n")
                detail_count += 1

    audit_headers = list(audit_rows[0])
    args.audit_json.write_text(
        json.dumps(
            {
                "headers": audit_headers,
                "rows": [[row[header] for header in audit_headers] for row in audit_rows],
                "summary": {
                    "sample_count": len(sample_rows),
                    "unique_patient_count": len({patient_id(row["ID"].strip()) for row in sample_rows}),
                    "selected_sample_visit_count": selected_visit_total,
                    "detail_row_count": detail_count,
                    "rule_counts": dict(rule_counts),
                    "visit_anchor_definition": "同一患者同一住院次数内最早申请时间",
                },
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    print(json.dumps({"detail_rows": detail_count, "audit_rows": len(audit_rows), "rules": dict(rule_counts)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
