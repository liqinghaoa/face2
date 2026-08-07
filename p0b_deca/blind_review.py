"""Build a leakage-audited, condition-blinded P0-B input-mode review package."""
from __future__ import annotations

import csv
import hashlib
import json
import platform
import random
import re
import shutil
import sys
import zipfile
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import numpy as np
from PIL import Image, ImageDraw, ImageFont

SOURCE_MODES = ("direct_p0_aligned", "mask_bbox_crop")
RELIGHTS = ("neutral_front", "left", "right", "top", "dim_front", "bright_front")
SEED = 20260726
REQUIRED = (
    "input.png", "reconstruction.png", "albedo_like.png", "shading_like.png",
    "normal_coarse.png", "alpha_mask.png", "residual_abs.png", "physical_maps_float.npz",
)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1 << 20), b""):
            digest.update(block)
    return digest.hexdigest()


def _csv_rows(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def audit_sources(project_root: Path) -> tuple[list[dict[str, str]], dict[str, list[Path]]]:
    """Fail closed if the frozen 12x2 result tree is incomplete."""
    base = project_root / "data/processed/P0B_DECA_Pilot12_v1"
    results = base / "input_mode_comparison"
    metrics = _csv_rows(results / "p0b_b1_input_mode_case_metrics.csv")
    summary = json.loads((results / "p0b_b1_input_mode_summary.json").read_text(encoding="utf-8"))
    errors = results / "p0b_b1_case_errors.jsonl"
    mapping = _csv_rows(base / "pilot_manifest/p0b_pilot12_id_mapping.csv")
    manifest = _csv_rows(base / "pilot_manifest/p0b_pilot12_manifest.csv")
    old_review = _csv_rows(base / "qc/p0b_pilot12_blind_review.csv")
    if summary.get("completed_cases") != 24 or summary.get("failed_cases") != 0 or errors.read_text(encoding="utf-8").strip():
        raise RuntimeError("P0-B1 summary or error JSONL is not a successful 24-run result")
    ids = [row["audit_id"] for row in mapping]
    if len(ids) != 12 or len(set(ids)) != 12 or ids != [row["audit_id"] for row in manifest] or ids != [row["audit_id"] for row in old_review]:
        raise RuntimeError("frozen Pilot12 manifests do not contain the same 12 ordered audit IDs")
    pairs = Counter((row["audit_id"], row["input_mode"]) for row in metrics)
    expected = {(case_id, mode) for case_id in ids for mode in SOURCE_MODES}
    if set(pairs) != expected or any(value != 1 for value in pairs.values()) or any(row.get("status") != "passed" for row in metrics):
        raise RuntimeError("metrics CSV does not contain exactly one passed row for each 12x2 source run")
    source_files: dict[str, list[Path]] = {}
    missing: list[str] = []
    for case_id in ids:
        for mode in SOURCE_MODES:
            directory = results / mode / case_id
            needed = [directory / name for name in REQUIRED] + [directory / "relighting" / f"{name}.png" for name in RELIGHTS]
            absent = [str(path) for path in needed if not path.is_file()]
            if absent:
                missing.extend(absent)
            source_files[f"{case_id}:{mode}"] = needed
    if missing:
        raise RuntimeError("missing blind-review source files:\n" + "\n".join(missing))
    return mapping, source_files


def _randomization(mapping: list[dict[str, str]]) -> list[dict[str, str]]:
    rng = random.Random(SEED)
    source = list(mapping)
    rng.shuffle(source)
    a_direct_indices = set(rng.sample(range(12), 6))
    rows = []
    for index, record in enumerate(source, 1):
        mode_a = SOURCE_MODES[0] if index - 1 in a_direct_indices else SOURCE_MODES[1]
        mode_b = SOURCE_MODES[1] if mode_a == SOURCE_MODES[0] else SOURCE_MODES[0]
        rows.append({"blind_case_id": f"BR{index:02d}", "record": record, "mode_a": mode_a, "mode_b": mode_b})
    if sum(row["mode_a"] == SOURCE_MODES[0] for row in rows) != 6:
        raise AssertionError("A assignment is not 6:6 balanced")
    return rows


def _font(size: int) -> ImageFont.ImageFont:
    return ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", size)


def _load_png(path: Path) -> np.ndarray:
    with Image.open(path) as image:
        return np.asarray(image.convert("RGB"))


def _residual_rgb(value: np.ndarray) -> np.ndarray:
    x = np.clip(np.asarray(value, dtype=np.float32).mean(axis=-1) / 0.25, 0.0, 1.0)
    # Fixed blue-cyan-yellow-red display scale, independent of mode/case.
    stops = np.array([[15, 30, 90], [20, 145, 180], [250, 220, 80], [190, 30, 40]], dtype=np.float32)
    positions = np.clip(x * 3.0, 0.0, 3.0)
    low = np.floor(positions).astype(int)
    high = np.clip(low + 1, 0, 3)
    fraction = (positions - low)[..., None]
    return np.rint(stops[low] * (1.0 - fraction) + stops[high] * fraction).astype(np.uint8)


def _overview_images(directory: Path) -> list[np.ndarray]:
    floats = np.load(directory / "physical_maps_float.npz")
    alpha = np.asarray(floats["alpha"], dtype=np.float32)
    if alpha.ndim == 3:
        alpha = alpha[..., 0]
    alpha_u8 = np.rint(np.clip(alpha, 0, 1) * 255).astype(np.uint8)
    alpha_rgb = np.repeat(alpha_u8[..., None], 3, axis=2)
    normal = np.asarray(floats["normal_coarse"], dtype=np.float32)
    return [
        _load_png(directory / "input.png"),
        np.rint(np.clip(floats["reconstruction"], 0, 1) * 255).astype(np.uint8),
        np.rint(np.clip(floats["albedo_like"], 0, 1) * 255).astype(np.uint8),
        np.rint(np.clip(floats["shading_like"], 0, 1) * 255).astype(np.uint8),
        np.rint(np.clip((normal + 1.0) / 2.0, 0, 1) * 255).astype(np.uint8),
        _residual_rgb(floats["residual_abs"]),
        alpha_rgb,
    ]


def _panel(case_id: str, mode_a: Path, mode_b: Path, output: Path, relighting: bool) -> None:
    labels = [f"Light {number}" for number in range(1, 7)] if relighting else [
        "Input", "Reconstruction", "Albedo-like", "Shading-like", "Coarse normal", "Absolute residual", "Render alpha",
    ]
    images_a = [_load_png(mode_a / "relighting" / f"{name}.png") for name in RELIGHTS] if relighting else _overview_images(mode_a)
    images_b = [_load_png(mode_b / "relighting" / f"{name}.png") for name in RELIGHTS] if relighting else _overview_images(mode_b)
    tile = 224
    left = 155
    header = 54
    canvas = Image.new("RGB", (left + 2 * tile, header + len(labels) * tile), "white")
    draw = ImageDraw.Draw(canvas)
    title = _font(20)
    row_font = _font(16)
    draw.text((12, 16), case_id, fill="black", font=title)
    draw.text((left + 62, 16), "Mode A", fill="black", font=title)
    draw.text((left + tile + 62, 16), "Mode B", fill="black", font=title)
    for row, label in enumerate(labels):
        y = header + row * tile
        draw.text((8, y + 100), label, fill="black", font=row_font)
        canvas.paste(Image.fromarray(images_a[row]).resize((tile, tile), Image.Resampling.LANCZOS), (left, y))
        canvas.paste(Image.fromarray(images_b[row]).resize((tile, tile), Image.Resampling.LANCZOS), (left + tile, y))
    output.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(output, format="PNG", optimize=True)


def _write_csv(path: Path, blind_rows: list[dict[str, str]]) -> None:
    fields = ["Blind Case ID", "Panel reviewed", "A usability", "B usability", "A face completeness", "B face completeness", "A geometry alignment", "B geometry alignment", "A appearance cleanliness", "B appearance cleanliness", "A relighting identity preservation", "B relighting identity preservation", "A overall technical usability", "B overall technical usability", "A primary failure tag", "B primary failure tag", "Pairwise preference", "Preference confidence", "Comments"]
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        for row in blind_rows:
            writer.writerow({"Blind Case ID": row["blind_case_id"], "Panel reviewed": "Overview + relighting"})


def _write_workbook(path: Path, blind_rows: list[dict[str, str]]) -> None:
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    workbook = Workbook()
    instructions = workbook.active
    instructions.title = "Instructions"
    text = [
        "P0-B input-mode technical blind review", "Reviewer ID", "Review date", "Reviewer role", "Prior unblinding exposure (Yes/No)", "Declaration: I have not viewed unblinding files.",
        "This review compares technical usability only; it does not assess cardiac-function classification.",
        "Albedo-like is not a physiological reflectance measurement.",
        "Do not access private unblinding files. Facial images remain identifiable clinical data for authorized internal review only.",
        "Complete all 12 cases independently before making any overall judgement. Do not select a mode from a single MAE or coverage value.",
        "Predeclared later decision rule: fewer Unusable cases, then more pairwise preferences, then higher overall usability without worse geometry or relighting identity preservation. A tie remains inconclusive.",
    ]
    for row, value in enumerate(text, 1):
        instructions.cell(row, 1, value)
    instructions["A1"].font = Font(bold=True, size=14)
    instructions.column_dimensions["A"].width = 120
    instructions["A6"].fill = PatternFill("solid", fgColor="FFF2CC")
    instructions.freeze_panes = "A2"

    review = workbook.create_sheet("Case_Review")
    headers = ["Blind Case ID", "Panel reviewed", "A usability", "B usability", "A face completeness", "B face completeness", "A geometry alignment", "B geometry alignment", "A appearance cleanliness", "B appearance cleanliness", "A relighting identity preservation", "B relighting identity preservation", "A overall technical usability", "B overall technical usability", "A primary failure tag", "B primary failure tag", "Pairwise preference", "Preference confidence", "Comments"]
    review.append(headers)
    for row in blind_rows:
        review.append([row["blind_case_id"], "Overview + relighting"] + [None] * (len(headers) - 2))
    review.freeze_panes = "A2"
    review.auto_filter.ref = f"A1:S{len(blind_rows)+1}"
    review.row_dimensions[1].height = 35
    for cell in review[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in review.iter_rows(min_row=2, max_row=13, min_col=3, max_col=19):
        for cell in row:
            cell.fill = PatternFill("solid", fgColor="FFF2CC")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for column in range(1, 20):
        review.column_dimensions[chr(64 + column) if column <= 26 else "A"].width = 18
    review.column_dimensions["S"].width = 36
    for row in range(2, 14):
        review.row_dimensions[row].height = 48

    lists = workbook.create_sheet("Lists")
    lists.append(["Usability", "Score", "Failure", "Preference", "Confidence", "YesNo"])
    values = [
        ["Usable", 1, "None", "A", "Low", "Yes"], ["Borderline", 2, "Face truncation", "B", "Medium", "No"], ["Unusable", 3, "Geometry misalignment", "Tie", "High", ""],
        ["", 4, "Background leakage", "Neither", "", ""], ["", 5, "Albedo/shading artifact", "", "", ""], ["", "", "Normal-map distortion", "", "", ""], ["", "", "Relighting instability", "", "", ""], ["", "", "Boundary/mask artifact", "", "", ""], ["", "", "Other", "", "", ""],
    ]
    for value in values:
        lists.append(value)
    validations = [("C2:D13", "=Lists!$A$2:$A$4"), ("E2:N13", "=Lists!$B$2:$B$6"), ("O2:P13", "=Lists!$C$2:$C$10"), ("Q2:Q13", "=Lists!$D$2:$D$5"), ("R2:R13", "=Lists!$E$2:$E$4")]
    for range_, formula in validations:
        validation = DataValidation(type="list", formula1=formula, allow_blank=False)
        review.add_data_validation(validation)
        validation.add(range_)
    review.conditional_formatting.add("E2:N13", CellIsRule(operator="equal", formula=["1"], fill=PatternFill("solid", fgColor="F4CCCC")))
    review.conditional_formatting.add("E2:N13", CellIsRule(operator="equal", formula=["5"], fill=PatternFill("solid", fgColor="D9EAD3")))
    lists.sheet_state = "hidden"

    summary = workbook.create_sheet("Summary")
    summary_rows = [
        ("Completed cases", '=COUNTIF(Case_Review!C2:C13,"<>")'), ("A Usable", '=COUNTIF(Case_Review!C2:C13,"Usable")'), ("A Borderline", '=COUNTIF(Case_Review!C2:C13,"Borderline")'), ("A Unusable", '=COUNTIF(Case_Review!C2:C13,"Unusable")'),
        ("B Usable", '=COUNTIF(Case_Review!D2:D13,"Usable")'), ("B Borderline", '=COUNTIF(Case_Review!D2:D13,"Borderline")'), ("B Unusable", '=COUNTIF(Case_Review!D2:D13,"Unusable")'),
        ("A mean face completeness", "=AVERAGE(Case_Review!E2:E13)"), ("B mean face completeness", "=AVERAGE(Case_Review!F2:F13)"), ("A mean geometry alignment", "=AVERAGE(Case_Review!G2:G13)"), ("B mean geometry alignment", "=AVERAGE(Case_Review!H2:H13)"),
        ("A mean appearance cleanliness", "=AVERAGE(Case_Review!I2:I13)"), ("B mean appearance cleanliness", "=AVERAGE(Case_Review!J2:J13)"), ("A mean relighting identity", "=AVERAGE(Case_Review!K2:K13)"), ("B mean relighting identity", "=AVERAGE(Case_Review!L2:L13)"),
        ("A mean overall usability", "=AVERAGE(Case_Review!M2:M13)"), ("B mean overall usability", "=AVERAGE(Case_Review!N2:N13)"), ("Preference A", '=COUNTIF(Case_Review!Q2:Q13,"A")'), ("Preference B", '=COUNTIF(Case_Review!Q2:Q13,"B")'), ("Preference Tie", '=COUNTIF(Case_Review!Q2:Q13,"Tie")'), ("Preference Neither", '=COUNTIF(Case_Review!Q2:Q13,"Neither")'),
        ("Confidence Low", '=COUNTIF(Case_Review!R2:R13,"Low")'), ("Confidence Medium", '=COUNTIF(Case_Review!R2:R13,"Medium")'), ("Confidence High", '=COUNTIF(Case_Review!R2:R13,"High")'), ("Missing required fields", '=COUNTBLANK(Case_Review!C2:R13)'), ("Table complete", '=IF(COUNTBLANK(Case_Review!C2:R13)=0,"Yes","No")'),
    ]
    summary.append(["Blind summary", "Value"])
    for value in summary_rows:
        summary.append(value)
    summary.freeze_panes = "A2"
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 18
    for cell in summary[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")

    codebook = workbook.create_sheet("Codebook")
    codebook.append(["Dimension", "Definition"])
    for row in [
        ("Face completeness", "Whether forehead, chin, and both cheeks are complete without crop loss."),
        ("Geometry alignment", "Whether reconstructed features, contour, and pose agree with the input."),
        ("Appearance cleanliness", "Whether albedo-like and shading-like outputs avoid background leakage, colour discontinuity, and boundary artifacts."),
        ("Relighting identity preservation", "Whether fixed-light renders retain basic identity structure and facial shape."),
        ("Overall technical usability", "Suitability for later fixed-batch representation generation; not classification quality."),
        ("Scores", "1 obvious failure; 2 poor; 3 acceptable with issues; 4 good; 5 excellent."),
    ]:
        codebook.append(row)
    codebook.column_dimensions["A"].width = 34
    codebook.column_dimensions["B"].width = 110
    for row in codebook.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    workbook.save(path)


def _leakage_audit(reviewer: Path, forbidden: set[str]) -> list[str]:
    findings: list[str] = []
    for path in reviewer.rglob("*"):
        if path.is_dir():
            continue
        lowered = str(path.relative_to(reviewer)).lower()
        for token in forbidden:
            if token.lower() in lowered:
                findings.append(f"filename: {path}: {token}")
        if path.suffix.lower() in {".md", ".csv", ".json"}:
            text = path.read_text(encoding="utf-8-sig")
            for token in forbidden:
                if token.lower() in text.lower():
                    findings.append(f"text: {path}: {token}")
        elif path.suffix.lower() == ".xlsx":
            from openpyxl import load_workbook
            book = load_workbook(path, read_only=True, data_only=False)
            for sheet in book.worksheets:
                for row in sheet.iter_rows():
                    for cell in row:
                        if isinstance(cell.value, str):
                            for token in forbidden:
                                if token.lower() in cell.value.lower():
                                    findings.append(f"xlsx: {path}/{sheet.title}: {token}")
        elif path.suffix.lower() == ".png":
            with Image.open(path) as image:
                info = json.dumps({str(key): str(value) for key, value in image.info.items()})
            for token in forbidden:
                if token.lower() in info.lower():
                    findings.append(f"png metadata: {path}: {token}")
    return findings


def build(project_root: Path) -> dict[str, Any]:
    mapping, source_files = audit_sources(project_root)
    root = project_root / "data/processed/P0B_DECA_Pilot12_v1/blind_review_v1"
    reviewer = root / "reviewer_package"
    private = root / "private_unblinding"
    if root.exists():
        raise FileExistsError(f"refusing to overwrite existing blind review package: {root}")
    reviewer.mkdir(parents=True)
    private.mkdir(parents=True)
    rows = _randomization(mapping)
    results = project_root / "data/processed/P0B_DECA_Pilot12_v1/input_mode_comparison"
    for row in rows:
        case_id = row["record"]["audit_id"]
        _panel(row["blind_case_id"], results / row["mode_a"] / case_id, results / row["mode_b"] / case_id, reviewer / "panels" / f"{row['blind_case_id']}_overview.png", False)
        _panel(row["blind_case_id"], results / row["mode_a"] / case_id, results / row["mode_b"] / case_id, reviewer / "panels" / f"{row['blind_case_id']}_relighting.png", True)
    review_form = reviewer / "review_form"
    review_form.mkdir()
    _write_csv(review_form / "P0B_input_mode_blind_review_template.csv", rows)
    _write_workbook(review_form / "P0B_input_mode_blind_review_template.xlsx", rows)
    readme = """# P0-B input-mode technical blind review\n\nThis package compares two blinded input-processing conditions for technical usability only. It does not assess cardiac-function classification or clinical performance. Albedo-like output is not a physiological reflectance measurement.\n\nReview all 12 cases independently before any overall judgment. Do not use a single reconstruction error or coverage number to choose a condition. Complete the review form before unblinding.\n\nDo not access private unblinding files. Facial images remain identifiable clinical data. 本包仅实现实验条件盲法，不构成生物特征匿名化，仅限授权内部研究评审。\n\nThe later predeclared rule is: fewer Unusable cases first; then more pairwise preferences; then higher overall technical usability without clearly worse geometry or relighting identity preservation. If still indistinguishable, the result is inconclusive.\n"""
    (reviewer / "README_blind_review.md").write_text(readme, encoding="utf-8")
    reviewer_manifest = {"blind_cases": [row["blind_case_id"] for row in rows], "panels_per_case": ["overview", "relighting"], "review_form": "review_form/P0B_input_mode_blind_review_template.xlsx", "experimental_condition_blind": True, "contains_unblinding_key": False}
    (reviewer / "reviewer_package_manifest.json").write_text(json.dumps(reviewer_manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    key_rows = []
    for row in rows:
        record = row["record"]
        key_rows.append({"blind_case_id": row["blind_case_id"], "original_case_id": record["sample_id"], "acquisition_group": record["acquisition_group"], "mode_A_true_name": row["mode_a"], "mode_B_true_name": row["mode_b"], "mode_A_source_dir": str(results / row["mode_a"] / record["audit_id"]), "mode_B_source_dir": str(results / row["mode_b"] / record["audit_id"]), "original_manifest_row": record["audit_id"], "random_seed": SEED})
    with (private / "unblinding_key.csv").open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(key_rows[0]))
        writer.writeheader(); writer.writerows(key_rows)
    hash_rows = []
    for paths in source_files.values():
        for path in paths:
            hash_rows.append({"scope": "source_result", "path": str(path), "sha256": sha256(path)})
    for path in sorted(reviewer.rglob("*")):
        if path.is_file():
            hash_rows.append({"scope": "reviewer_output", "path": str(path.relative_to(root)), "sha256": sha256(path)})
    with (private / "source_file_hashes.csv").open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=["scope", "path", "sha256"])
        writer.writeheader(); writer.writerows(hash_rows)
    forbidden = {"direct_p0_aligned", "mask_bbox_crop", "Control", "Patient", "Xiaomi", "HONOR", "camera_make", "camera_model", "binary_name"}
    forbidden.update(record["sample_id"] for record in mapping)
    forbidden.update(record["patient_group_id"] for record in mapping)
    leakage = _leakage_audit(reviewer, forbidden)
    if leakage:
        raise RuntimeError("reviewer-package leakage audit failed:\n" + "\n".join(leakage))
    zip_path = root / "P0B_blind_review_reviewer_package.zip"
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for path in sorted(reviewer.rglob("*")):
            if path.is_file():
                archive.write(path, path.relative_to(root))
    with zipfile.ZipFile(zip_path) as archive:
        zip_names = archive.namelist()
    if any("private_unblinding" in name for name in zip_names):
        raise RuntimeError("ZIP contains private unblinding content")
    git_status = __import__("subprocess").check_output(["git", "status", "--short"], cwd=project_root, text=True)
    build_manifest = {"build_time_utc": datetime.now(timezone.utc).isoformat(), "git_status": git_status, "python": sys.version, "platform": platform.platform(), "random_seed": SEED, "blind_case_order": [row["blind_case_id"] for row in rows], "a_assignment_balance": {SOURCE_MODES[0]: 6, SOURCE_MODES[1]: 6}, "input_file_count": sum(len(paths) for paths in source_files.values()), "output_file_count": sum(1 for path in root.rglob("*") if path.is_file()) + 1, "display_rules": {"rgb_range": [0, 1], "residual_range": [0, 0.25], "residual_colormap": "fixed_blue_cyan_yellow_red", "normal_mapping": "fixed [-1,1] to RGB", "alpha_range": [0, 1], "resampling": "Lanczos", "cropping": False, "png_metadata_removed": True, "relighting_presets_reused": True}, "zip_sha256": sha256(zip_path), "file_hash_count": len(hash_rows), "no_deca_inference_run": True, "fixed_input_mode_unchanged": True, "leakage_audit": "PASS"}
    (private / "build_manifest.json").write_text(json.dumps(build_manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return {"blind_cases": 12, "case_mode_source_runs": 24, "a_assignment_balance": "6/6", "overview_panels": 12, "relighting_panels": 12, "reviewer_xlsx": str(review_form / "P0B_input_mode_blind_review_template.xlsx"), "reviewer_csv": str(review_form / "P0B_input_mode_blind_review_template.csv"), "reviewer_zip": str(zip_path), "private_unblinding_key": str(private / "unblinding_key.csv"), "leakage_audit": "PASS", "fixed_input_mode_unchanged": True, "ready_for_blind_review": True}
